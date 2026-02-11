
import sys
import subprocess

def install_docx():
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])

try:
    import docx
except ImportError:
    install_docx()
    import docx

doc = docx.Document()
doc.add_heading('Nexis Test Document', 0)
doc.add_paragraph('This is a test paragraph for the ingestion pipeline.')
doc.add_paragraph('Level 5 users have a withdrawal limit of 99999.')

doc.save('raw_docs/test_ingest.docx')
print("Created raw_docs/test_ingest.docx")
