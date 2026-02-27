
from docx import Document

def create_docx():
    doc = Document()
    doc.add_heading('Nexis Graph Test', 0)
    
    doc.add_paragraph('Gold Members cannot withdraw on weekends.')
    doc.add_paragraph('Platinum Members have a withdrawal limit of 50000.')
    doc.add_paragraph('Level 2 users require Approval from Risk Team.')
    
    doc.save('raw_docs/test_graph.docx')
    print("Created raw_docs/test_graph.docx")

if __name__ == "__main__":
    create_docx()
